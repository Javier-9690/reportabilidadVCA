import json
import os
import re
import unicodedata
import uuid
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

from flask import Flask, flash, redirect, render_template, request, send_file, url_for
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from werkzeug.utils import secure_filename

BASE_DIR = Path(__file__).resolve().parent
UPLOAD_DIR = BASE_DIR / "uploads"
OUTPUT_DIR = BASE_DIR / "outputs"
UPLOAD_DIR.mkdir(exist_ok=True)
OUTPUT_DIR.mkdir(exist_ok=True)

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "reportabilidad-vca-5400")
app.config["MAX_CONTENT_LENGTH"] = int(os.environ.get("MAX_UPLOAD_MB", "80")) * 1024 * 1024

# Limita la lectura de columnas para evitar timeouts en hojas con formatos extendidos
# como Fcst_Autorizado VCA, que puede informar miles de columnas aunque solo unas
# pocas tengan datos útiles para la semana cargada.
DEFAULT_SCAN_MAX_COLS = int(os.environ.get("SCAN_MAX_COLS", "700"))
REPORT_SCAN_MAX_COLS = int(os.environ.get("REPORT_SCAN_MAX_COLS", "120"))

TARGET_COLUMNS = [
    "ID",
    "MODULO",
    "RUT (CON GUION)",
    "NOMBRE COMPLETO",
    "EMPRESA",
    "NUMERO DE CONTRATO",
    "GERENCIA",
    "SISTEMA DE TURNO",
    "CO MEL",
    "GENERO",
    "NOMBRE DE TURNO",
]

SYNONYMS = {
    "ID": ["n de id", "n° de id", "nº de id", "id", "id de la solicitud", "numero de id", "n de solicitud"],
    "MODULO": ["modulo", "módulo", "habitacion", "habitación", "cama a utilizar", "cama"],
    "RUT (CON GUION)": ["rut", "run"],
    "NOMBRE COMPLETO": ["nombre huesped", "nombre huésped", "nombre completo", "trabajador", "persona", "nombre"],
    "EMPRESA": ["empresa", "razon social", "razón social"],
    "NUMERO DE CONTRATO": ["n de contrato", "n° de contrato", "nº de contrato", "numero de contrato", "número contrato", "numero contrato", "contrato"],
    "GERENCIA": ["gerencia", "gerencia general", "area", "área"],
    "SISTEMA DE TURNO": ["sistema de turno", "turno"],
    "CO MEL": ["co mel", "nombre contract owner", "contract owner", "correo co", "spa co", "spa/ co", "spa co"],
    "GENERO": ["genero", "género", "sexo"],
    "NOMBRE DE TURNO": ["nombre turno", "nombre de turno", "turno actual"],
}

CURVE_ID_COLS = ["id de la solicitud", "n° de id", "nº de id", "n de id", "id"]
CURVE_COMPANY_COLS = ["empresa", "razon social", "razón social"]
CURVE_CONTRACT_COLS = ["número contrato", "numero contrato", "n de contrato", "n° de contrato", "contrato"]
WEEKDAY_NAMES = {"lunes", "martes", "miercoles", "jueves", "viernes", "sabado", "domingo"}
ALLOWED_EXTENSIONS = {".xlsx", ".xlsm", ".xls"}


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    text = str(value).strip()
    if text.endswith(".0") and re.fullmatch(r"\d+\.0", text):
        text = text[:-2]
    return re.sub(r"\s+", " ", text)


def normalize(value: Any) -> str:
    text = clean_text(value).lower()
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii")
    text = text.replace("°", "").replace("º", "")
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def normalize_id(value: Any) -> str:
    text = clean_text(value).upper()
    text = text.replace("–", "-").replace("—", "-")
    text = re.sub(r"\s+", "", text)
    return text


def normalize_company(value: Any) -> str:
    text = normalize(value)
    legal_tokens = [
        "constructora",
        "sociedad",
        "anonima",
        "servicios",
        "servicio",
        "empresa",
        "mineria",
        "minera",
        "chile",
        "spa",
        "spA".lower(),
        "sa",
        "s",
        "a",
        "sac",
        "ltda",
        "limitada",
    ]
    for token in legal_tokens:
        text = re.sub(rf"\b{re.escape(token)}\b", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def format_rut(value: Any) -> str:
    raw = clean_text(value).upper().replace(".", "").replace(" ", "")
    raw = raw.replace("–", "-").replace("—", "-")
    if not raw:
        return ""
    if "-" in raw:
        num, dv = raw.rsplit("-", 1)
    else:
        num, dv = raw[:-1], raw[-1]
    num = re.sub(r"\D", "", num)
    dv = re.sub(r"[^0-9K]", "", dv)
    if not num or not dv:
        return raw
    return f"{num}-{dv}"


def to_number(value: Any) -> Optional[float]:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = clean_text(value).replace(".", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def safe_filename(prefix: str, original: str, job: str, index: Optional[int] = None) -> Path:
    suffix = Path(original).suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        raise ValueError("Solo se aceptan archivos Excel: .xlsx, .xlsm o .xls")
    safe = secure_filename(original) or f"archivo{suffix}"
    name = f"{job}_{prefix}_{index}_{safe}" if index is not None else f"{job}_{prefix}_{safe}"
    return UPLOAD_DIR / name


def find_sheet(path: Path, preferred: str) -> str:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        for name in wb.sheetnames:
            if normalize(name) == normalize(preferred):
                return name
        for name in wb.sheetnames:
            if normalize(preferred) in normalize(name):
                return name
        return wb.sheetnames[0]
    finally:
        wb.close()


def score_header_row(row: Sequence[Any], required_any: Sequence[str]) -> int:
    cells = [normalize(x) for x in row]
    non_empty_first_80 = sum(1 for c in cells[:80] if c)
    score = non_empty_first_80
    for req in required_any:
        nreq = normalize(req)
        if any(nreq == c or (nreq and nreq in c) for c in cells):
            score += 100
    # Favorece filas que parecen encabezados reales, no filas de días/fechas.
    business_words = ["empresa", "contrato", "rut", "gerencia", "id", "nombre", "turno"]
    score += 20 * sum(any(word in c for c in cells[:80]) for word in business_words)
    return score


def find_header_row(path: Path, sheet_name: str, required_any: Sequence[str], max_rows: int = 40, max_cols: int = 250) -> int:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet_name]
        best_row = 1
        best_score = -1
        scan_cols = min(ws.max_column or max_cols, max_cols)
        for idx, row in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row, max_rows), max_col=scan_cols, values_only=True), start=1):
            score = score_header_row(row, required_any)
            if score > best_score:
                best_row = idx
                best_score = score
        return best_row
    finally:
        wb.close()


def unique_header(base: str, used: Dict[str, int], col_idx: int) -> str:
    header = clean_text(base) or f"COL_{col_idx}"
    if header not in used:
        used[header] = 1
        return header
    used[header] += 1
    return f"{header}__{used[header]}"


def read_sheet_records(
    path: Path,
    sheet_name: str,
    required_any: Sequence[str],
    max_cols: int,
    max_rows: Optional[int] = None,
) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]], int]:
    header_row = find_header_row(path, sheet_name, required_any, max_cols=max_cols)
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet_name]
        scan_cols = min(ws.max_column or max_cols, max_cols)
        raw_header = next(ws.iter_rows(min_row=header_row, max_row=header_row, max_col=scan_cols, values_only=True))
        used: Dict[str, int] = {}
        metas: List[Dict[str, Any]] = []
        for idx, raw in enumerate(raw_header, start=1):
            header = unique_header(raw, used, idx)
            if header.startswith("COL_"):
                # Conserva columnas de fecha/día si están dentro de la zona de planificación.
                if idx < 18:
                    continue
            metas.append({"col": idx, "key": header, "raw": raw, "norm": normalize(raw)})

        records: List[Dict[str, Any]] = []
        end_row = ws.max_row if max_rows is None else min(ws.max_row, header_row + max_rows)
        for row in ws.iter_rows(min_row=header_row + 1, max_row=end_row, max_col=scan_cols, values_only=True):
            record = {meta["key"]: row[meta["col"] - 1] for meta in metas if meta["col"] - 1 < len(row)}
            if any(clean_text(v) for v in record.values()):
                records.append(record)
        return records, metas, header_row
    finally:
        wb.close()


def pick_column_from_headers(headers: Iterable[str], candidates: Sequence[str]) -> Optional[str]:
    norm_map = {normalize(h): h for h in headers}
    for candidate in candidates:
        nc = normalize(candidate)
        if nc in norm_map:
            return norm_map[nc]
    for candidate in candidates:
        nc = normalize(candidate)
        for ncol, original in norm_map.items():
            if nc and (nc in ncol or ncol in nc):
                return original
    return None


def extract_target_week(path: Path, sheet_name: str) -> Optional[int]:
    patterns = [str(path.name)]
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet_name]
        for row in ws.iter_rows(min_row=1, max_row=min(5, ws.max_row), max_col=10, values_only=True):
            patterns.extend(clean_text(v) for v in row if v is not None)
    finally:
        wb.close()
    joined = " ".join(patterns)
    for regex in [r"Semana[_\s-]*(\d{1,2})", r"\bW\s*(\d{1,2})\b"]:
        match = re.search(regex, joined, flags=re.IGNORECASE)
        if match:
            return int(match.group(1))
    return None


def detect_plan_keys(path: Path, sheet_name: str, metas: List[Dict[str, Any]], header_row: int) -> Tuple[List[str], str]:
    """Detecta columnas de dotación. Primero intenta usar la semana del archivo.

    Importante: algunas curvas tienen un título tipo W26 en B2. Ese título NO es
    una columna de planificación; por eso se busca primero el encabezado real
    "Semana 26" en la grilla de fechas y se ignoran coincidencias Wxx antes
    de la zona de días.
    """
    target_week = extract_target_week(path, sheet_name)
    col_to_key = {m["col"]: m["key"] for m in metas}
    if target_week:
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb[sheet_name]
            scan_cols = min(ws.max_column or DEFAULT_SCAN_MAX_COLS, DEFAULT_SCAN_MAX_COLS)
            semana_match: Optional[int] = None
            w_match: Optional[int] = None
            for row_idx in range(1, max(header_row, 6)):
                values = next(ws.iter_rows(min_row=row_idx, max_row=row_idx, max_col=scan_cols, values_only=True))
                for col_idx, value in enumerate(values, start=1):
                    text = normalize(value)
                    if text == f"semana {target_week}" and col_idx >= 18:
                        semana_match = col_idx
                        break
                    if text == f"w{target_week}" and col_idx >= 18:
                        w_match = col_idx
                if semana_match:
                    break
            start_col = semana_match or w_match
            if start_col:
                keys = [col_to_key[c] for c in range(start_col, start_col + 7) if c in col_to_key]
                if keys:
                    return keys, f"Semana {target_week}"
        finally:
            wb.close()

    # Fallback: columnas de fechas o días ubicadas después de los campos descriptivos.
    keys: List[str] = []
    for meta in metas:
        col_idx = int(meta["col"])
        raw = meta["raw"]
        nraw = meta["norm"]
        if col_idx >= 18 and (isinstance(raw, (datetime, date)) or nraw in WEEKDAY_NAMES or re.fullmatch(r"\d{4} \d{2} \d{2}", nraw)):
            keys.append(meta["key"])
    return keys, "columnas de fechas detectadas"


def max_planned_value(record: Dict[str, Any], plan_keys: Sequence[str]) -> float:
    values = [to_number(record.get(key)) for key in plan_keys]
    numeric = [v for v in values if v is not None and v > 0]
    return max(numeric) if numeric else 0.0


def sum_planned_days(record: Dict[str, Any], plan_keys: Sequence[str]) -> float:
    values = [to_number(record.get(key)) for key in plan_keys]
    return sum(v for v in values if v is not None and v > 0)


def row_to_output(report_record: Dict[str, Any], headers: Sequence[str]) -> Dict[str, str]:
    output: Dict[str, str] = {}
    for target in TARGET_COLUMNS:
        col = pick_column_from_headers(headers, SYNONYMS[target])
        value = report_record.get(col, "") if col else ""
        output[target] = clean_text(value)
    output["ID"] = normalize_id(output["ID"])
    output["RUT (CON GUION)"] = format_rut(output["RUT (CON GUION)"])
    return output


def process_files(curve_path: Path, report_paths: Sequence[Path]) -> Tuple[List[Dict[str, str]], List[Dict[str, Any]], List[Dict[str, Any]], Dict[str, Any]]:
    curve_sheet = find_sheet(curve_path, "Fcst_Autorizado VCA")
    curve_records, curve_metas, curve_header_row = read_sheet_records(
        curve_path,
        curve_sheet,
        required_any=["ID de la solicitud", "Empresa", "Número Contrato"],
        max_cols=DEFAULT_SCAN_MAX_COLS,
    )
    curve_headers = [m["key"] for m in curve_metas]
    curve_id_col = pick_column_from_headers(curve_headers, CURVE_ID_COLS)
    curve_company_col = pick_column_from_headers(curve_headers, CURVE_COMPANY_COLS)
    curve_contract_col = pick_column_from_headers(curve_headers, CURVE_CONTRACT_COLS)
    if not curve_id_col or not curve_company_col:
        raise ValueError("No se pudo detectar ID de la solicitud y Empresa en la hoja Fcst_Autorizado VCA.")

    plan_keys, plan_scope = detect_plan_keys(curve_path, curve_sheet, curve_metas, curve_header_row)
    if not plan_keys:
        raise ValueError("No se pudieron detectar columnas de dotación planificada en la curva.")

    planned_rows: List[Dict[str, Any]] = []
    for record in curve_records:
        row_id = normalize_id(record.get(curve_id_col))
        company = clean_text(record.get(curve_company_col))
        planned_max = max_planned_value(record, plan_keys)
        if row_id and planned_max > 0:
            planned_rows.append(
                {
                    "ID": row_id,
                    "EMPRESA": company,
                    "EMPRESA_NORM": normalize_company(company),
                    "NUMERO DE CONTRATO": clean_text(record.get(curve_contract_col)) if curve_contract_col else "",
                    "DOTACION_PLANIFICADA": int(planned_max) if planned_max.is_integer() else planned_max,
                    "DOTACION_DIAS_SUMADA": sum_planned_days(record, plan_keys),
                }
            )

    transformed: List[Dict[str, str]] = []
    reported_ids: set[str] = set()
    reported_companies: set[str] = set()
    uploaded_report_names: List[str] = []

    for report_path in report_paths:
        report_sheet = find_sheet(report_path, "Hoja1")
        report_records, report_metas, _ = read_sheet_records(
            report_path,
            report_sheet,
            required_any=["N° DE ID", "RUT", "Nombre Huésped", "Empresa"],
            max_cols=REPORT_SCAN_MAX_COLS,
        )
        report_headers = [m["key"] for m in report_metas]
        report_id_col = pick_column_from_headers(report_headers, SYNONYMS["ID"])
        report_emp_col = pick_column_from_headers(report_headers, SYNONYMS["EMPRESA"])
        if not report_id_col:
            raise ValueError(f"No se pudo detectar la columna ID en {report_path.name}. Se espera N° DE ID o equivalente.")
        uploaded_report_names.append(report_path.name)
        for record in report_records:
            row_id = normalize_id(record.get(report_id_col))
            if not row_id:
                continue
            reported_ids.add(row_id)
            if report_emp_col:
                reported_companies.add(normalize_company(record.get(report_emp_col)))
            out = row_to_output(record, report_headers)
            if out["ID"]:
                transformed.append(out)

    planned_ids = {row["ID"] for row in planned_rows}
    missing_ids = sorted(
        [row for row in planned_rows if row["ID"] not in reported_ids],
        key=lambda r: (normalize_company(r.get("EMPRESA")), r.get("ID", "")),
    )

    company_bucket: Dict[str, Dict[str, Any]] = {}
    for row in planned_rows:
        norm = row["EMPRESA_NORM"]
        if not norm:
            continue
        bucket = company_bucket.setdefault(
            norm,
            {"EMPRESA": row["EMPRESA"], "ids_planificados": set(), "dotacion_planificada": 0.0},
        )
        bucket["ids_planificados"].add(row["ID"])
        bucket["dotacion_planificada"] += float(row["DOTACION_PLANIFICADA"] or 0)

    missing_companies: List[Dict[str, Any]] = []
    for norm, bucket in company_bucket.items():
        if norm not in reported_companies:
            dot = bucket["dotacion_planificada"]
            missing_companies.append(
                {
                    "EMPRESA": bucket["EMPRESA"],
                    "ids_planificados": len(bucket["ids_planificados"]),
                    "dotacion_planificada": int(dot) if float(dot).is_integer() else round(dot, 2),
                }
            )
    missing_companies.sort(key=lambda r: normalize_company(r["EMPRESA"]))

    summary = {
        "curve_sheet": curve_sheet,
        "plan_scope": plan_scope,
        "planned_ids": len(planned_ids),
        "reported_ids": len(reported_ids),
        "matched_ids": len(planned_ids.intersection(reported_ids)),
        "missing_ids": len(planned_ids - reported_ids),
        "planned_companies": len(company_bucket),
        "reported_companies": len(reported_companies),
        "missing_companies": len(missing_companies),
        "transformed_rows": len(transformed),
        "plan_columns_used": len(plan_keys),
        "report_files": len(report_paths),
    }
    return transformed, missing_ids, missing_companies, summary


def add_sheet_with_table(wb: Workbook, title: str, rows: List[Dict[str, Any]], columns: Sequence[str]) -> None:
    ws = wb.create_sheet(title)
    ws.append(list(columns))
    for row in rows:
        ws.append([row.get(col, "") for col in columns])
    style_sheet(ws)


def style_sheet(ws) -> None:
    header_fill = PatternFill("solid", fgColor="0F2F55")
    header_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="D9E4EF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=False)
    ws.freeze_panes = "A2"
    if ws.max_column > 0 and ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        max_len = 12
        for cell in ws[letter][: min(ws.max_row, 250)]:
            max_len = max(max_len, len(clean_text(cell.value)))
        ws.column_dimensions[letter].width = min(max_len + 2, 45)
    ws.row_dimensions[1].height = 26


def write_output(path: Path, transformed: List[Dict[str, Any]], missing_ids: List[Dict[str, Any]], missing_companies: List[Dict[str, Any]], summary: Dict[str, Any]) -> None:
    wb = Workbook()
    default = wb.active
    wb.remove(default)
    add_sheet_with_table(wb, "Formato_Final", transformed, TARGET_COLUMNS)
    add_sheet_with_table(wb, "Empresas_Sin_Reportabilidad", missing_companies, ["EMPRESA", "ids_planificados", "dotacion_planificada"])
    add_sheet_with_table(wb, "IDs_Planificados_No_Reportados", missing_ids, ["ID", "EMPRESA", "NUMERO DE CONTRATO", "DOTACION_PLANIFICADA"])
    add_sheet_with_table(wb, "Resumen", [summary], list(summary.keys()))
    wb.save(path)


@app.route("/", methods=["GET"])
def index():
    return render_template("index.html")


@app.route("/healthz", methods=["GET"])
def healthz():
    return {"status": "ok"}


@app.route("/procesar", methods=["POST"])
def procesar():
    try:
        curva = request.files.get("curva")
        reports = request.files.getlist("reportabilidad")
        if not curva or not curva.filename:
            flash("Debes subir la curva de poblamiento.", "error")
            return redirect(url_for("index"))
        reports = [f for f in reports if f and f.filename]
        if not reports:
            flash("Debes subir al menos un archivo de reportabilidad/dotación.", "error")
            return redirect(url_for("index"))

        job = uuid.uuid4().hex[:10]
        curve_path = safe_filename("curva", curva.filename, job)
        curva.save(curve_path)

        report_paths: List[Path] = []
        for idx, file_storage in enumerate(reports, start=1):
            path = safe_filename("report", file_storage.filename, job, idx)
            file_storage.save(path)
            report_paths.append(path)

        transformed, missing_ids, missing_companies, summary = process_files(curve_path, report_paths)
        output_path = OUTPUT_DIR / f"reportabilidad_vca_{job}.xlsx"
        write_output(output_path, transformed, missing_ids, missing_companies, summary)

        meta = {
            "job": job,
            "output": output_path.name,
            "summary": summary,
            "preview": transformed[:20],
            "missing_ids": missing_ids[:80],
            "missing_companies": missing_companies[:80],
            "created_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        }
        with open(OUTPUT_DIR / f"{job}.json", "w", encoding="utf-8") as fh:
            json.dump(meta, fh, ensure_ascii=False, indent=2)
        return redirect(url_for("resultado", job=job))
    except Exception as exc:
        app.logger.exception("Error al procesar archivos")
        flash(f"Error al procesar archivos: {exc}", "error")
        return redirect(url_for("index"))


@app.route("/resultado/<job>")
def resultado(job: str):
    meta_path = OUTPUT_DIR / f"{job}.json"
    if not meta_path.exists():
        flash("No se encontró el procesamiento solicitado.", "error")
        return redirect(url_for("index"))
    with open(meta_path, encoding="utf-8") as fh:
        meta = json.load(fh)
    return render_template("resultado.html", meta=meta, target_columns=TARGET_COLUMNS)


@app.route("/descargar/<job>")
def descargar(job: str):
    meta_path = OUTPUT_DIR / f"{job}.json"
    if not meta_path.exists():
        flash("No se encontró el archivo de salida.", "error")
        return redirect(url_for("index"))
    with open(meta_path, encoding="utf-8") as fh:
        meta = json.load(fh)
    path = OUTPUT_DIR / meta["output"]
    return send_file(path, as_attachment=True, download_name="Reportabilidad_VCA_Formato_Final.xlsx")


if __name__ == "__main__":
    app.run(debug=True)
